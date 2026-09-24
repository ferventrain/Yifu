"""Post-run QC for brain-mode vessel samples: human-checkable visuals plus a
minimal machine verdict.

Outputs under <sample>/qc/:
  registration_views.png  - three orthogonal mid-planes of the (coarse)
                            registration NIfTI, per-view normalized, with the
                            registered label contour (red) and the hemisphere
                            split line (yellow, axial+coronal views).
  seg_blocks/block_NN.zarr- N random full-resolution blocks, dataset '0' =
                            raw signal, dataset '1' = segmentation mask.
  verdict.json            - four hard checks only (label magnitude, hemisphere
                            balance, 25 um grid shape, registration
                            correlation). overall FAIL must stop the queue.

Designed to be cheap: reads the ~250 MB NIfTI, mid-plane slices of the label
Zarr, the small warped-atlas TIFF stack, and a handful of 64x512x512 blocks —
never a full-resolution volume scan.
"""
from __future__ import annotations

import argparse
import json
import logging
import random
import sys
import time
from datetime import datetime
from pathlib import Path

import numpy as np

PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from pipeline_modules.utils.run_manifest import write_run_manifest
from pipeline_modules.utils.zarr_io import create_array, open_output_group, open_zarr_array

logger = logging.getLogger(__name__)

N_SEG_BLOCKS = 5
BLOCK_ZYX = (64, 512, 512)
RANDOM_SEED = 20260924

LABEL_VOXEL_RANGE = (5e9, 5e11)
HEMISPHERE_BALANCE_RANGE = (0.35, 0.65)
GRID_TOLERANCE = 0.20
# Calibrated on MPTP_2 (known-good): pearson_r = 0.437 — template-vs-
# autofluorescence is a cross-modality comparison, so genuine registrations
# land around 0.4; the historical failure mode (empty warp) is caught by the
# warped_nonzero count instead. Misaligned warps trend to ~0 or negative.
CORRELATION_FAIL = 0.25
CORRELATION_WARN = 0.35


def _norm_u8(plane: np.ndarray) -> np.ndarray:
    lo, hi = np.percentile(plane, [1.0, 99.8])
    if hi <= lo:
        hi = lo + 1
    return np.clip((plane.astype(np.float32) - lo) / (hi - lo) * 255, 0, 255).astype(np.uint8)


def _label_contour(presence: np.ndarray) -> np.ndarray:
    from scipy import ndimage

    if not presence.any():
        return np.zeros(presence.shape, dtype=bool)
    return ndimage.binary_dilation(presence, iterations=2) & ~presence


def _read_root_stats(results_dir: Path) -> dict | None:
    import openpyxl

    paths = sorted(results_dir.glob("*_brain_distribution_stats.xlsx"))
    if not paths:
        return None
    workbook = openpyxl.load_workbook(paths[0], read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(min_row=1, max_row=2, values_only=True))
    workbook.close()
    if len(rows) < 2:
        return None
    return dict(zip(rows[0], rows[1]))


def check_root_stats(root: dict | None) -> dict:
    if root is None:
        return {"root_stats": {"status": "FAIL", "reason": "results xlsx not found or unreadable"}}
    left = int(root.get("Left Total Voxels") or 0)
    right = int(root.get("Right Total Voxels") or 0)
    total = left + right or int(root.get("Total Voxels") or 0)
    magnitude_status = "PASS" if LABEL_VOXEL_RANGE[0] <= total <= LABEL_VOXEL_RANGE[1] else "FAIL"
    balance = left / (left + right) if (left + right) else 0.0
    balance_status = "PASS" if HEMISPHERE_BALANCE_RANGE[0] <= balance <= HEMISPHERE_BALANCE_RANGE[1] else "FAIL"
    return {
        "label_total_voxels": {"status": magnitude_status, "value": total, "range": list(LABEL_VOXEL_RANGE)},
        "hemisphere_balance": {"status": balance_status, "value": round(balance, 4), "left": left, "right": right},
    }


def check_grid_shape(nii_path: Path, original_shape_path: Path, resolution_xyz, target_xyz) -> dict:
    import nibabel as nib

    if not nii_path.exists():
        return {"status": "FAIL", "reason": f"missing {nii_path}"}
    nii_shape_xyz = tuple(int(v) for v in nib.load(str(nii_path)).shape)  # x,y,z
    try:
        original = json.loads(original_shape_path.read_text(encoding="utf-8"))["original_shape"]  # z,y,x
    except (OSError, KeyError, json.JSONDecodeError):
        return {"status": "WARN", "reason": f"missing/unreadable {original_shape_path}"}
    expected_zyx = [o * r / t for o, r, t in zip(original, resolution_xyz[::-1], target_xyz[::-1])]
    expected_xyz = expected_zyx[::-1]
    ratios = [actual / exp if exp else 0 for actual, exp in zip(nii_shape_xyz, expected_xyz)]
    if any(abs(ratio - 1.0) > GRID_TOLERANCE for ratio in ratios):
        return {"status": "FAIL", "ratios_xyz": [round(r, 3) for r in ratios],
                "nii_shape_xyz": nii_shape_xyz, "expected_xyz": [round(e, 1) for e in expected_xyz]}
    return {"status": "PASS", "ratios_xyz": [round(r, 3) for r in ratios]}


def check_registration_correlation(nii_path: Path, warped_dir: Path) -> dict:
    import nibabel as nib
    import tifffile
    from scipy import ndimage

    if not nii_path.exists() or not warped_dir.is_dir():
        return {"status": "FAIL", "reason": "missing nii or warped-atlas dir"}
    nii_zyx = np.transpose(np.asanyarray(nib.load(str(nii_path)).dataobj), (2, 1, 0))  # -> z,y,x
    slices = sorted(warped_dir.glob("*.tif*"))
    if not slices:
        return {"status": "FAIL", "reason": "no warped atlas slices"}
    warped = np.stack([tifffile.imread(str(s)) for s in slices]).astype(np.float32)
    del slices
    resampled = ndimage.zoom(nii_zyx.astype(np.float32), [
        w / n for w, n in zip(warped.shape, nii_zyx.shape)], order=1, mode="nearest")
    mask = warped > 0
    if mask.sum() < 1000:
        return {"status": "FAIL", "reason": "warped atlas is (near-)empty", "warped_nonzero": int(mask.sum())}
    a = resampled[mask]
    b = warped[mask]
    if a.std() < 1e-6 or b.std() < 1e-6:
        return {"status": "FAIL", "reason": "degenerate variance"}
    corr = float(np.corrcoef(a, b)[0, 1])
    status = "PASS" if corr >= CORRELATION_WARN else ("WARN" if corr >= CORRELATION_FAIL else "FAIL")
    return {"status": status, "pearson_r": round(corr, 4), "warped_nonzero": int(mask.sum())}


def _match_plane(presence: np.ndarray, plane: np.ndarray) -> np.ndarray:
    """Crop/pad a label-derived plane to the display plane shape."""
    if presence.shape == plane.shape:
        return presence
    out = np.zeros(plane.shape, dtype=bool)
    h, w = min(presence.shape[0], plane.shape[0]), min(presence.shape[1], plane.shape[1])
    out[:h, :w] = presence[:h, :w]
    return out


def render_registration_views(nii_path: Path, label_zarr_path: Path, hemi_zarr_path: Path, out_png: Path) -> None:
    import nibabel as nib
    from PIL import Image, ImageDraw

    nii_xyz = np.asanyarray(nib.load(str(nii_path)).dataobj)  # x,y,z
    nx, ny, nz = nii_xyz.shape
    label = open_zarr_array(label_zarr_path) if label_zarr_path.exists() else None  # z,y,x
    hemi = open_zarr_array(hemi_zarr_path) if hemi_zarr_path.exists() else None
    lz, ly, lx = (int(v) for v in label.shape) if label is not None else (0, 0, 0)
    fx, fy, fz = lx / max(nx, 1), ly / max(ny, 1), lz / max(nz, 1)
    sx = [max(int(round(v)), 1) for v in (fz, fy, fx)]  # strides on label z,y,x

    # Display convention: rows = first listed axis, cols = second.
    # axial (rows=y, cols=x), coronal (rows=z, cols=x), sagittal (rows=z, cols=y).
    planes = {
        "axial": _norm_u8(np.transpose(nii_xyz[:, :, nz // 2])),          # -> (y, x)
        "coronal": _norm_u8(np.transpose(nii_xyz[:, ny // 2, :])),        # -> (z, x)
        "sagittal": _norm_u8(np.transpose(nii_xyz[nx // 2, :, :])),       # -> (z, y)
    }
    label_planes: dict[str, np.ndarray | None] = {"axial": None, "coronal": None, "sagittal": None}
    hemi_plane: np.ndarray | None = None
    if label is not None:
        label_planes["axial"] = np.asarray(label[lz // 2])[:: sx[1], :: sx[2]] > 0          # (y, x)
        label_planes["coronal"] = np.asarray(label[:, ly // 2, :])[:: sx[0], :: sx[2]] > 0  # (z, x)
        label_planes["sagittal"] = np.asarray(label[:, :, lx // 2])[:: sx[0], :: sx[1]] > 0  # (z, y)
    if hemi is not None:
        hemi_plane = np.asarray(hemi[lz // 2])[:: sx[1], :: sx[2]]  # (y, x): 0 bg, 1 left, 2 right

    import zarr as _zarr

    try:
        split_x = dict(_zarr.open(str(hemi_zarr_path), mode="r").attrs).get("split_x")
    except Exception:
        split_x = None

    images = []
    for name, plane in planes.items():
        rgb = np.repeat(plane[..., None], 3, axis=2)
        presence = label_planes.get(name)
        if presence is not None:
            rgb[_label_contour(_match_plane(presence, plane))] = (255, 60, 60)
        img = Image.fromarray(rgb)
        draw = ImageDraw.Draw(img)
        if name in ("axial", "coronal") and split_x:
            # cols are the x axis on both views; scale full-res split into nii x.
            draw.line([(int(split_x / fx), 0), (int(split_x / fx), img.height)], fill=(255, 220, 0), width=2)
        draw.text((6, 4), name, fill=(255, 255, 255))
        images.append(img)

    # Hemisphere-verification panel: axial plane tinted left=red / right=blue
    # on the autofluorescence background — shows whether the split plane
    # actually separates the two lobes.
    if hemi_plane is not None:
        rgb = np.repeat(planes["axial"][..., None], 3, axis=2)
        matched = _match_plane(hemi_plane, planes["axial"])
        rgb[matched == 1] = (255, 90, 90)
        rgb[matched == 2] = (90, 120, 255)
        img = Image.fromarray(rgb)
        draw = ImageDraw.Draw(img)
        draw.text((6, 4), "axial L/R check", fill=(255, 255, 255))
        images.append(img)

    height = max(img.height for img in images)
    sheet = Image.new("RGB", (sum(img.width for img in images) + 16 * (len(images) - 1), height), (20, 20, 20))
    offset = 0
    for img in images:
        sheet.paste(img, (offset, 0))
        offset += img.width + 16
    sheet.save(out_png)


def save_segmentation_blocks(signal_zarr_path: Path, mask_zarr_path: Path, label_zarr_path: Path,
                             out_dir: Path, sample_name: str) -> list[Path]:
    signal = open_zarr_array(signal_zarr_path)
    mask = open_zarr_array(mask_zarr_path)
    label = open_zarr_array(label_zarr_path)
    nz, ny, nx = (int(v) for v in signal.shape)
    bz, by, bx = (min(b, n) for b, n in zip(BLOCK_ZYX, (nz, ny, nx)))
    rng = random.Random(RANDOM_SEED)
    label_center = (label is not None)
    written = []
    attempts = 0
    while len(written) < N_SEG_BLOCKS and attempts < 40:
        attempts += 1
        z0 = rng.randrange(0, nz - bz + 1)
        y0 = rng.randrange(0, ny - by + 1)
        x0 = rng.randrange(0, nx - bx + 1)
        if label_center and not np.any(np.asarray(label[z0:z0 + bz, y0:y0 + by, x0:x0 + bx]) > 0):
            continue
        mask_block = np.asarray(mask[z0:z0 + bz, y0:y0 + by, x0:x0 + bx])
        if not np.any(mask_block > 0):
            continue
        signal_block = np.asarray(signal[z0:z0 + bz, y0:y0 + by, x0:x0 + bx])
        block_path = out_dir / f"block_{len(written):02d}.zarr"
        group = open_output_group(block_path, overwrite=True)
        create_array(group, "0", shape=signal_block.shape, chunks=(min(32, bz), by, bx),
                     dtype=signal_block.dtype, data=signal_block)
        create_array(group, "1", shape=mask_block.shape, chunks=(min(32, bz), by, bx),
                     dtype=mask_block.dtype, data=mask_block)
        group.attrs["block_offset_zyx"] = [z0, y0, x0]
        group.attrs["block_shape_zyx"] = [bz, by, bx]
        group.attrs["sample"] = sample_name
        group.attrs["channels"] = {"0": "raw signal", "1": "segmentation mask"}
        written.append(block_path)
        logger.info("saved QC block %s at zyx=%s", block_path.name, (z0, y0, x0))
    if len(written) < N_SEG_BLOCKS:
        logger.warning("only %d/%d QC blocks saved (empty mask or label)", len(written), N_SEG_BLOCKS)
    return written


def run_qc(
    sample_dir: Path,
    *,
    config: dict,
    nii_path: Path | None = None,
    label_zarr_path: Path | None = None,
    hemi_zarr_path: Path | None = None,
    warped_dir: Path | None = None,
    results_dir: Path | None = None,
    signal_zarr_path: Path | None = None,
    mask_zarr_path: Path | None = None,
    output_dir: Path | None = None,
) -> dict:
    sample_dir = Path(sample_dir)
    signal_ch = str(config["input"]["channels"]["signal"])
    nii_path = nii_path or sample_dir / "ch0_downsample" / "volume.nii.gz"
    label_zarr_path = label_zarr_path or sample_dir / "upsampled_atlas_label.zarr"
    hemi_zarr_path = hemi_zarr_path or sample_dir / "atlas_label_hemisphere.zarr"
    warped_dir = warped_dir or sample_dir / "ch0_warped_image"
    results_dir = results_dir or sample_dir / "results"
    signal_zarr_path = signal_zarr_path or sample_dir / f"ch{signal_ch}.zarr"
    mask_zarr_path = mask_zarr_path or sample_dir / f"ch{signal_ch}_mask.zarr"
    output_dir = output_dir or sample_dir / "qc"
    output_dir.mkdir(parents=True, exist_ok=True)

    checks: dict[str, dict] = {}
    root = _read_root_stats(results_dir)
    checks.update(check_root_stats(root))
    checks["grid_shape"] = check_grid_shape(
        nii_path, sample_dir / "original_shape.json",
        [float(v) for v in config["input"]["resolution_xyz"]],
        [float(v) for v in config["preprocessing"]["downsample"]["target_resolution_xyz"]],
    )
    checks["registration_correlation"] = check_registration_correlation(nii_path, warped_dir)

    try:
        render_registration_views(nii_path, label_zarr_path, hemi_zarr_path, output_dir / "registration_views.png")
    except Exception:
        logger.exception("registration views failed")

    try:
        save_segmentation_blocks(signal_zarr_path, mask_zarr_path, label_zarr_path,
                                 output_dir / "seg_blocks", sample_dir.name)
    except Exception:
        logger.exception("segmentation blocks failed")

    split_x_attr = None
    try:
        import zarr

        split_x_attr = dict(zarr.open(str(hemi_zarr_path), mode="r").attrs).get("split_x")
    except Exception:
        logger.warning("hemisphere zarr attrs unreadable: %s", hemi_zarr_path)

    statuses = [str(check.get("status")) for check in checks.values()]
    overall = "FAIL" if "FAIL" in statuses else ("WARN" if "WARN" in statuses else "PASS")
    verdict = {
        "sample": sample_dir.name,
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "checks": checks,
        "hemisphere_split_x": split_x_attr,
        "overall": overall,
    }
    (output_dir / "verdict.json").write_text(json.dumps(verdict, indent=1, ensure_ascii=False), encoding="utf-8")
    logger.info("QC verdict: %s (%s)", overall, json.dumps({k: v.get("status") for k, v in checks.items()}))
    return verdict


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Post-run QC visuals + minimal verdict")
    parser.add_argument("--sample_dir", required=True)
    parser.add_argument("--config", default="", help="config.json (defaults to sample_dir/config.json)")
    parser.add_argument("--nii", default="", help="Override registration NIfTI path")
    parser.add_argument("--label_zarr", default="")
    parser.add_argument("--hemi_zarr", default="")
    parser.add_argument("--warped_dir", default="")
    parser.add_argument("--results_dir", default="")
    parser.add_argument("--signal_zarr", default="")
    parser.add_argument("--mask_zarr", default="")
    parser.add_argument("--output_dir", default="")
    return parser.parse_args()


def main() -> int:
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(name)s: %(message)s")
    args = parse_args()
    sample_dir = Path(args.sample_dir)
    config_path = Path(args.config) if args.config else sample_dir / "config.json"
    config = json.loads(config_path.read_text(encoding="utf-8"))
    started = time.time()
    verdict = run_qc(
        sample_dir,
        config=config,
        nii_path=Path(args.nii) if args.nii else None,
        label_zarr_path=Path(args.label_zarr) if args.label_zarr else None,
        hemi_zarr_path=Path(args.hemi_zarr) if args.hemi_zarr else None,
        warped_dir=Path(args.warped_dir) if args.warped_dir else None,
        results_dir=Path(args.results_dir) if args.results_dir else None,
        signal_zarr_path=Path(args.signal_zarr) if args.signal_zarr else None,
        mask_zarr_path=Path(args.mask_zarr) if args.mask_zarr else None,
        output_dir=Path(args.output_dir) if args.output_dir else None,
    )
    write_run_manifest(
        sample_dir / "qc",
        module="visualization.pipeline_qc",
        entrypoint="run_qc",
        inputs={"sample_dir": str(sample_dir), "config": str(config_path)},
        outputs=[sample_dir / "qc" / "verdict.json", sample_dir / "qc" / "registration_views.png"],
        started_at=started,
        extra={"overall": verdict["overall"]},
    )
    return 0 if verdict["overall"] != "FAIL" else 1


if __name__ == "__main__":
    raise SystemExit(main())
